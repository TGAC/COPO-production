import json
import os
import re
import time

import openpyxl
import pytest
from playwright.sync_api import Page, TimeoutError as PlaywrightTimeoutError, expect

from test.playwright.e2e.helpers import (
    EDP_PROFILE_FIELDS,
    create_edp_profile,
    delete_profile_by_title,
    dismiss_tour_if_present,
    edp_jira_ticket_number,
    poll_sample_accession,
    read_sapio_project,
    run_django_shell,
    upload_manifest,
)

PROFILE_TITLE = 'Playwright EDP LIMS Submission Journey'
PROFILE_DESCRIPTION = 'A description that is at least twenty characters long.'

TESTFILES_DIR = os.path.normpath(os.path.join(os.path.dirname(__file__), '..', '..', 'testfiles'))
SAMPLE_MANIFEST = os.path.join(TESTFILES_DIR, 'COPO_FAANG_manifest_success1.xlsx')
# Holds only what a customer fills in. Everything EI/Sapio generates per
# project (Project ID, budget holder, JIRA ticket, Sample ID, EI sample name,
# plate and well) changes every run, so the test downloads the manifest COPO
# pre-fills for this run's profile and copies these answers into it.
EDP_LIMS_MANIFEST = os.path.join(TESTFILES_DIR, 'edp_lims_test_manifest_tissue_success1.xlsx')

FAANG_SAMPLE_CHECKLIST = 'COPO_FAANG'
# The EI_EDP schema's checklists_filter maps sample type 'Tissue' to this
# checklist, and the LIMS page offers only that one for a Tissue profile.
EDP_CHECKLIST = 'version_tissue'

# The manifest's study-level answers, keyed by their EI_EDP schema label.
STUDY_CELLS = {'Email Address': 'E6', 'Sample Return': 'L8', 'Health and Safety': 'L9'}
# The sample sheet's EI-only block (Sample ID ... Project ID) ends where this
# column starts; every column from here rightwards is the customer's.
FIRST_SUBMITTER_LABEL = 'Submitter Sample Reference'
# These three are VLOOKUP array formulas into the sample_metadata sheet,
# which COPO fills from the profile's own ENA samples. COPO reads the upload
# with data_only=True (EDPSchemasSpreadsheet.loadManifest), i.e. it reads
# Excel's cached formula results -- and a workbook saved by openpyxl has none,
# so the formulas would read back empty. Resolve them here instead.
SAMPLE_METADATA_LOOKUPS = {
    'NCBI Taxon ID': 'taxon_id',
    'Scientific Name (Latin)': 'scientific_name',
    'BioSample Accession': 'biosampleAccession',
}


def _header_columns(sheet, row):
    return {str(cell.value).strip(): cell.column for cell in sheet[row] if cell.value not in (None, '')}


def _find_header_row(sheet):
    return next(row[0].row for row in sheet.iter_rows() if row[0].value == 'Sample ID')


def _build_edp_manifest(downloaded_path, patched_path, accession):
    # Returns (sapio_sample_id, expected) where expected maps each schema
    # label to the value this manifest submits for it.
    downloaded = openpyxl.load_workbook(downloaded_path)
    answers = openpyxl.load_workbook(EDP_LIMS_MANIFEST)

    sample_sheet, answer_sheet = downloaded['sample'], answers['sample']
    expected = {}
    for label, cell in STUDY_CELLS.items():
        sample_sheet[cell] = expected[label] = answer_sheet[cell].value

    header_row = _find_header_row(sample_sheet)
    answer_header_row = _find_header_row(answer_sheet)
    columns = _header_columns(sample_sheet, header_row)
    answer_columns = _header_columns(answer_sheet, answer_header_row)
    data_row, answer_row = header_row + 1, answer_header_row + 1

    sapio_sample_id = sample_sheet.cell(data_row, columns['Sample ID']).value
    assert sapio_sample_id, 'Downloaded manifest has no pre-filled Sapio Sample ID'
    assert sample_sheet.cell(data_row + 1, columns['Sample ID']).value in (None, ''), (
        'Downloaded manifest has more than one sample row; the answers file describes one'
    )

    reference = answer_sheet.cell(answer_row, answer_columns[FIRST_SUBMITTER_LABEL]).value
    metadata_sheet = downloaded['sample_metadata']
    metadata_header = [cell.value for cell in metadata_sheet[1]]
    metadata = {
        values[0]: dict(zip(metadata_header, values))
        for values in metadata_sheet.iter_rows(min_row=2, values_only=True)
        if values[0]
    }
    assert reference in metadata, (
        f'Sample {reference!r} is missing from the downloaded sample_metadata sheet '
        f"(found {sorted(metadata)}); the ENA sample's name must match the manifest's "
        f'Submitter Sample Reference'
    )
    assert metadata[reference]['biosampleAccession'] == accession

    first_submitter_column = columns[FIRST_SUBMITTER_LABEL]
    for label, column in columns.items():
        if column < first_submitter_column:
            continue
        if label in SAMPLE_METADATA_LOOKUPS:
            value = metadata[reference][SAMPLE_METADATA_LOOKUPS[label]]
        else:
            value = answer_sheet.cell(answer_row, answer_columns[label]).value
        sample_sheet.cell(data_row, column).value = expected[label] = value

    # The index sheet's Sample ID is the Earlham Institute's (EI's) (Sapio's) ID for the sample, which
    # only exists once the profile has been saved.
    index_sheet, answer_index_sheet = downloaded['index'], answers['index']
    index_columns = _header_columns(index_sheet, 1)
    for label, answer_column in _header_columns(answer_index_sheet, 1).items():
        if label in ('Project ID', 'Sample ID'):
            continue
        index_sheet.cell(2, index_columns[label]).value = answer_index_sheet.cell(2, answer_column).value
    index_sheet.cell(2, index_columns['Sample ID']).value = sapio_sample_id

    downloaded.save(patched_path)
    return sapio_sample_id, expected


def _sapio_field_mapping():
    # Which Sapio field each manifest label is written to, straight from the
    # EI_EDP schema's sapio_name column ('Object:Field') -- the same mapping
    # SapioAdapter._submit_manifest() uses. Only Project and Sample fields
    # that aren't 'protected' are written by the adapter.
    code = (
        'import json\n'
        'from src.apps.copo_single_cell_submission.utils.da import SinglecellSchemas\n'
        f"schema = SinglecellSchemas().get_schema(schema_name='EI_EDP', target_id={EDP_CHECKLIST!r})\n"
        'fields = [\n'
        "    {'label': str(f.get('term_label', '')).strip(), 'sapio_name': f.get('sapio_name'),\n"
        "     'behavior': f.get('term_manifest_behavior')}\n"
        "    for component in ('study', 'sample') for f in schema.get(component, [])\n"
        ']\n'
        "print('MAPPING_RESULT:' + json.dumps(fields, default=str))\n"
    )
    marker_line = next(line for line in run_django_shell(code).splitlines() if line.startswith('MAPPING_RESULT:'))
    mapping = {}
    for field in json.loads(marker_line[len('MAPPING_RESULT:'):]):
        sapio_name = field['sapio_name']
        if not isinstance(sapio_name, str) or ':' not in sapio_name or field['behavior'] == 'protected':
            continue
        sapio_object, sapio_field = sapio_name.split(':', 1)
        if sapio_object in ('Project', 'Sample') and sapio_field:
            mapping[field['label']] = (sapio_object, sapio_field)
    return mapping


def _same_value(expected, actual):
    # Sapio returns numbers as numbers and may differ in case for picklist
    # values; the manifest holds what the customer typed.
    try:
        return float(expected) == float(actual)
    except (TypeError, ValueError):
        return str(expected).strip().casefold() == str(actual).strip().casefold()


def _poll_sapio_study_status(profile_id, study_id, timeout=60, interval=5):
    # submit_edp_to_sapio() records the outcome on the singlecell study as
    # status_sapio / accession_sapio / error_sapio (update_component_status
    # suffixes each key with the repository).
    code = (
        'from src.apps.copo_single_cell_submission.utils.da import Singlecell\n'
        f"records = Singlecell(profile_id={profile_id!r}).execute_query({{'profile_id': {profile_id!r}}})\n"
        'study = {}\n'
        'for r in records:\n'
        "    for s in r.get('components', {}).get('study', []):\n"
        f"        if s.get('study_id') == {study_id!r}:\n"
        '            study = s\n'
        "print('POLL_RESULT:' + '|'.join(["
        "str(study.get('status_sapio', '')), str(study.get('accession_sapio', '')), str(study.get('error_sapio', ''))"
        ']))\n'
    )
    deadline = time.time() + timeout
    status = ''
    while time.time() < deadline:
        marker_line = next(
            (line for line in run_django_shell(code).splitlines() if line.startswith('POLL_RESULT:')),
            'POLL_RESULT:||',
        )
        status, accession, error = marker_line[len('POLL_RESULT:'):].split('|', 2)
        if status == 'accepted' and accession:
            return accession
        if status == 'rejected' or error:
            raise AssertionError(f'Sapio submission recorded as failed: status={status!r} error={error!r}')
        time.sleep(interval)
    raise AssertionError(f'Sapio submission status never recorded (last status={status!r})')


@pytest.mark.external
def test_edp_manifest_submission_to_sapio(page: Page, tmp_path, sapio_test_instance, ena_dev_sandbox):
    # The EDP journey: EDP profile (creates its Sapio project and samples)
    # -> FAANG sample manifest -> real 'Submit to ENA' against EBI's dev
    # sandbox -> wait for a real BioSample accession -> download the LIMS
    # manifest COPO pre-fills for this project (which now lists that sample
    # and its accession) -> fill in the customer's answers -> validate and
    # save it -> 'Submit to Sapio' -> read the project back from Sapio and
    # check every mapped manifest field landed.
    #
    # Clearing leftovers by title also deletes their Sapio projects (the EDP
    # post-delete hook), so an interrupted run doesn't leak projects.
    delete_profile_by_title(PROFILE_TITLE)

    jira_ticket_number = edp_jira_ticket_number()
    _, profile_id, project_id = create_edp_profile(
        page, PROFILE_TITLE, PROFILE_DESCRIPTION, jira_ticket_number=jira_ticket_number, **EDP_PROFILE_FIELDS
    )
    assert re.fullmatch(r'PR\d+', project_id), f'Unexpected Sapio project id {project_id!r}'

    # --- Samples: upload, save and submit to ENA ---
    page.goto(f'/copo/copo_sample/view/general/{profile_id}/general_sample')
    page.wait_for_load_state('networkidle')
    dismiss_tour_if_present(page)

    page.select_option('#checklist_id', FAANG_SAMPLE_CHECKLIST)
    page.click('.new-general-sample-spreadsheet-template')
    page.wait_for_load_state('networkidle')
    page.wait_for_timeout(1000)
    upload_manifest(page, SAMPLE_MANIFEST)
    try:
        expect(page.locator('#save_sample_button')).to_be_enabled(timeout=60000)
    except AssertionError as error:
        # parse_sample_spreadsheet answers a failed validation with an empty
        # 400; the reasons only reach the page via the websocket.
        raise AssertionError(f"Sample manifest rejected: {page.locator('#sample_info').inner_text()}") from error
    page.click('#save_sample_button')

    page.wait_for_timeout(1000)
    dismiss_tour_if_present(page)
    page.locator('#sample_table tbody tr').first.click()
    page.click(".copo-dt[data-action='submit_sample']")
    # A fresh test user has no ENA credentials of its own, so COPO offers
    # its default; with valid credentials the submission just proceeds.
    try:
        page.locator('#repo_cred_modal .cred-modal-default').wait_for(state='visible', timeout=5000)
    except PlaywrightTimeoutError:
        pass
    else:
        page.click('#repo_cred_modal .cred-modal-default')

    accession = poll_sample_accession(profile_id)

    # --- LIMS manifest: download the pre-filled one, fill it in ---
    response = page.request.get(
        f'/copo/copo_single_cell/download_initial_manifest/{profile_id}/EI_EDP/{EDP_CHECKLIST}'
    )
    assert response.ok, f'Manifest download failed: {response.status}'
    downloaded_manifest = tmp_path / 'edp_manifest_downloaded.xlsx'
    downloaded_manifest.write_bytes(response.body())

    # The project-level cells come from the profile and can't be edited in
    # the manifest, so they prove the profile -> manifest pre-fill works.
    project_sheet = openpyxl.load_workbook(downloaded_manifest)['sample']
    assert project_sheet['C5'].value == project_id
    assert project_sheet['C4'].value == EDP_PROFILE_FIELDS['budget_user']
    assert project_sheet['F5'].value == jira_ticket_number

    patched_manifest = tmp_path / 'edp_manifest_patched.xlsx'
    sapio_sample_id, expected = _build_edp_manifest(downloaded_manifest, patched_manifest, accession)

    # --- LIMS manifest: validate and save ---
    page.goto(f'/copo/copo_single_cell/view/{profile_id}/EI_EDP/lims_integration')
    page.wait_for_load_state('networkidle')
    dismiss_tour_if_present(page)

    page.select_option('#checklist_id', EDP_CHECKLIST)
    page.click('.new-singlecell-spreadsheet-template')
    page.wait_for_load_state('networkidle')
    page.wait_for_timeout(1000)
    with page.expect_response(lambda r: 'parse_singlecell_spreadsheet' in r.url, timeout=120000) as response_info:
        upload_manifest(page, str(patched_manifest))
    response = response_info.value
    assert response.status == 200, f'LIMS manifest rejected: {response.text()[:2000]}'

    expect(page.locator('#save_singlecell_button')).to_be_enabled(timeout=60000)
    page.click('#save_singlecell_button')
    page.wait_for_load_state('networkidle')

    # --- Submit to Sapio ---
    # Same table mechanics as the single-cell journey: the study tab's table
    # is `singlecell_table_study`, the toolbar is cloned into every tab (so
    # scope the click to the study table's own button), and the study tab
    # auto-selects its first row after each refresh -- clicking the row
    # ourselves would race that.
    page.wait_for_timeout(1000)
    dismiss_tour_if_present(page)
    page.wait_for_function(
        "() => $.fn.dataTable.isDataTable('#singlecell_table_study') && "
        "$('#singlecell_table_study').DataTable().rows({selected: true}).count() > 0",
        timeout=30000,
    )

    # The Sapio write happens inside this request (no Celery), so its own
    # feedback is the outcome. 'warning' means some manifest samples weren't
    # found in the Sapio project and were skipped -- a failure here.
    with page.expect_response(lambda r: 'copo_forms' in r.url, timeout=120000) as response_info:
        page.click(".copo-dt[data-action='submit_singlecell_sapio'][data-table='singlecell_table_study']")
    response = response_info.value
    assert response.status == 200, f'submit_singlecell_sapio request failed: {response.status} {response.text()[:500]}'
    feedback = response.json().get('action_feedback', {})
    assert feedback.get('status') == 'success', f'Submit to Sapio failed: {feedback}'

    record_id = _poll_sapio_study_status(profile_id, project_id)

    # --- Read the project back from Sapio ---
    mapping = _sapio_field_mapping()
    checks = [
        (label, sapio_object, sapio_field, expected[label])
        for label, (sapio_object, sapio_field) in mapping.items()
        if expected.get(label) not in (None, '')
    ]
    assert checks, "No manifest fields map to Sapio -- has the EI_EDP schema's sapio_name column changed?"

    sapio = read_sapio_project(
        project_id,
        project_fields=['RecordId', 'ProjectName'] + [f for _, o, f, _ in checks if o == 'Project'],
        sample_fields=['SampleId'] + [f for _, o, f, _ in checks if o == 'Sample'],
    )
    assert sapio['exists'], f'Sapio project {project_id} not found'
    assert str(sapio['project']['RecordId']) == record_id, "COPO's Sapio accession isn't the project's record id"
    assert sapio['project']['ProjectName'] == jira_ticket_number

    sapio_sample = next((s for s in sapio['samples'] if s['SampleId'] == sapio_sample_id), None)
    assert sapio_sample, f'Sample {sapio_sample_id} not found under Sapio project {project_id}'

    mismatches = []
    for label, sapio_object, sapio_field, expected_value in checks:
        actual = (sapio['project'] if sapio_object == 'Project' else sapio_sample)[sapio_field]
        if not _same_value(expected_value, actual):
            mismatches.append(f'{label} -> {sapio_object}:{sapio_field}: manifest {expected_value!r}, Sapio {actual!r}')
    assert not mismatches, "Manifest fields that didn't land in Sapio:\n" + '\n'.join(mismatches)

    delete_profile_by_title(PROFILE_TITLE)
    assert not read_sapio_project(project_id)['exists'], f'Sapio project {project_id} was not deleted'
